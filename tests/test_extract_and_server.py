from __future__ import annotations

import io

import anyio
import pytest

from gdrive_mcp import extract, server
from gdrive_mcp.tools import Tools
from test_gdrive_mcp import FakeApi, cfg, jpeg

def _mini_pdf() -> bytes:
    stream = b"BT /F1 12 Tf 20 100 Td (Hello PDF) Tj ET"
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    return out + b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)


MINI_PDF = _mini_pdf()


def _saved(obj) -> bytes:
    buf = io.BytesIO()
    obj.save(buf)
    return buf.getvalue()


def test_pdf_text():
    assert "Hello PDF" in extract.extract_text(MINI_PDF, extract.PDF)


def test_docx_text_and_tables():
    import docx

    d = docx.Document()
    d.add_paragraph("Rubrik")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "a", "b"
    text = extract.extract_text(_saved(d), extract.DOCX)
    assert "Rubrik" in text and "a\tb" in text


def test_pptx_text_and_notes():
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "Bild ett"
    s.notes_slide.notes_text_frame.text = "Talarnot"
    text = extract.extract_text(_saved(p), extract.PPTX)
    assert "--- slide 1 ---" in text and "Bild ett" in text and "[notes] Talarnot" in text


def test_xlsx_values_and_row_cap(monkeypatch):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Data"
    for i in range(10):
        wb.active.append([i, "x", None])
    data = _saved(wb)
    text = extract.extract_text(data, extract.XLSX)
    assert "--- sheet Data ---" in text and "9\tx" in text
    monkeypatch.setattr(extract, "MAX_XLSX_ROWS", 3)
    assert "[truncated after 3 rows]" in extract.extract_text(data, extract.XLSX)


def test_image_downscaled_and_transparency_kept():
    from PIL import Image

    data, fmt = extract.image_for_model(jpeg(3000, 1000))
    assert fmt == "jpeg" and max(Image.open(io.BytesIO(data)).size) == extract.MAX_IMAGE_SIDE
    _, fmt = extract.image_for_model(jpeg(10, 10, mode="RGBA", fmt="PNG"))
    assert fmt == "png"


# --- MCP server end to end (in memory) ------------------------------------------------------------
@pytest.fixture
def fake_server(monkeypatch, tmp_path):
    t = Tools(FakeApi(), cfg(), audit_path=tmp_path / "audit.jsonl")
    monkeypatch.setattr(server, "tools", lambda: t)
    return server.mcp


def _run(coro_fn):
    return anyio.run(coro_fn)


def test_server_lists_all_tools(fake_server):
    async def go():
        from mcp.client import Client

        async with Client(fake_server) as c:
            return await c.list_tools()

    names = {t.name for t in _run(go).tools}
    assert len(names) == 26 and "drive_semantic_search" in names and {"drive_read", "docs_edit", "sheets_append", "drive_share"} <= names


def test_server_sends_results_once(fake_server):
    """Structured output would duplicate every result and add output schemas to the tool list."""
    async def go():
        from mcp.client import Client

        async with Client(fake_server) as c:
            return (await c.list_tools()).tools, await c.call_tool("drive_get", {"file_id": "D1"})

    tools, res = _run(go)
    assert not [t.name for t in tools if t.output_schema]
    assert res.structured_content is None and len(res.content) == 1


def test_tool_annotations(fake_server):
    async def go():
        from mcp.client import Client

        async with Client(fake_server) as c:
            return {t.name: t.annotations for t in (await c.list_tools()).tools}

    ann = _run(go)
    assert all(a is not None and a.open_world_hint is False for a in ann.values())
    assert ann["drive_read"].read_only_hint and ann["drive_semantic_search"].read_only_hint
    assert ann["drive_create"].read_only_hint is False and ann["drive_create"].destructive_hint is False
    for name in ("docs_edit", "sheets_write", "sheets_edit", "drive_trash", "drive_move", "drive_update_text"):
        assert ann[name].destructive_hint is True, name


def test_server_returns_policy_errors_as_tool_errors(fake_server):
    async def go():
        from mcp.client import Client

        async with Client(fake_server) as c:
            ok = await c.call_tool("drive_get", {"file_id": "D1"})
            bad = await c.call_tool("drive_get", {"file_id": "OUTFILE"})
            img = await c.call_tool("drive_read", {"file_id": "PIC"})
            return ok, bad, img

    ok, bad, img = _run(go)
    assert not ok.is_error and "ChatGPT/Projects/Handoff" in ok.content[0].text
    assert bad.is_error and "outside the allowed folders" in bad.content[0].text
    assert img.content[0].type == "image" and "untrusted" in img.content[1].text


def test_server_share_refusal_reaches_agent(fake_server):
    async def go():
        from mcp.client import Client

        async with Client(fake_server) as c:
            return await c.call_tool("drive_share", {"file_id": "D1", "email": "x@evil.test", "role": "writer"})

    res = _run(go)
    assert res.is_error and "only allowed" in res.content[0].text


def test_json_body_sent_as_object_is_accepted(fake_server):
    """ChatGPT sends a JSON file's content as the object itself, not as a string."""
    async def go():
        from mcp.client import Client

        async with Client(fake_server) as c:
            return await c.call_tool(
                "drive_create",
                {"folder_id": "PROJ", "name": "data.json", "kind": "text",
                 "content": {"theses": [{"title": "Guld", "note": "åäö"}]}},
            )

    res = _run(go)
    assert not res.is_error, res.content[0].text
    assert server.body({"a": 1}) == '{\n  "a": 1\n}'
    assert server.body('{"a": 1}') == '{"a": 1}'
    assert server.body(None) is None
