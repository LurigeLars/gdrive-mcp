"""Text from PDF/docx/pptx/xlsx and model-sized images. Pure functions on bytes; no network."""
from __future__ import annotations

import io

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCUMENT_MIMES = {PDF, DOCX, PPTX, XLSX}
IMAGE_MIMES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
MAX_IMAGE_BYTES = 5 * 1024 * 1024
MAX_IMAGE_SIDE = 1568
MAX_XLSX_ROWS = 5000  # per workbook


def extract_text(data: bytes, mime: str) -> str:
    if mime == PDF:
        return _pdf(data)
    if mime == DOCX:
        return _docx(data)
    if mime == PPTX:
        return _pptx(data)
    if mime == XLSX:
        return _xlsx(data)
    raise ValueError(f"no text extractor for {mime}")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(f"--- page {i} ---\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1))


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t, table in enumerate(d.tables, 1):
        out.append(f"--- table {t} ---")
        out += ["\t".join(c.text for c in row.cells) for row in table.rows]
    return "\n".join(out)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        out.append(f"--- slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                out += ["\t".join(c.text for c in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"[notes] {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(out)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out, rows = [], 0
    try:
        for ws in wb.worksheets:
            out.append(f"--- sheet {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                if rows >= MAX_XLSX_ROWS:
                    out.append(f"[truncated after {MAX_XLSX_ROWS} rows]")
                    return "\n".join(out)
                out.append("\t".join("" if v is None else str(v) for v in row).rstrip("\t"))
                rows += 1
    finally:
        wb.close()
    return "\n".join(out)


def image_for_model(data: bytes) -> tuple[bytes, str]:
    """Downscale to at most MAX_IMAGE_SIDE px and re-encode (JPEG, or PNG when transparent)."""
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        img.load()
        frame = img.convert("RGBA") if img.mode in ("RGBA", "LA", "P") else img.convert("RGB")
    frame.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    buf = io.BytesIO()
    if frame.mode == "RGBA":
        frame.save(buf, "PNG", optimize=True)
        return buf.getvalue(), "png"
    frame.save(buf, "JPEG", quality=85)
    return buf.getvalue(), "jpeg"
